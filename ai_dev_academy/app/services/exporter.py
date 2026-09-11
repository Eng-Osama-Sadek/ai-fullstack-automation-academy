import io
import arabic_reshaper
from bidi.algorithm import get_display
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

def fix_arabic_text(text: str) -> str:
    reshaped_text = arabic_reshaper.reshape(text)
    return get_display(reshaped_text)

def export_to_docx(title: str, content: str) -> io.BytesIO:
    doc = Document()
    doc.add_heading(title, level=0)
    
    for line in content.split('\n'):
        doc.add_paragraph(line)
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def export_to_pdf(title: str, content: str) -> io.BytesIO:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    
    pdf.setFont("Helvetica", 14)
    pdf.drawString(100, 750, f"Title: {title}")
    
    y = 720
    for line in content.split('\n'):
        reshaped_line = fix_arabic_text(line) if any(ord(char) > 127 for char in line) else line
        pdf.drawString(50, y, reshaped_line[:80])
        y -= 20
        if y < 50:
            pdf.showPage()
            y = 750
            
    pdf.save()
    buffer.seek(0)
    return buffer